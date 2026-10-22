# -*- coding: utf-8 -*-
"""assess_run -- Lumerical run result checker (pure file parsing, no lumapi).

Scope: only traps that actually failed in past runs; the evidence field names the
reference file (../references) that explains each trap.
Verdicts: PASS / WARN / FAIL (overall = worst).

Usage (CLI):
    <python> assess_run.py PATH [PATH ...] [options]

    PATH may be
      *_p0.log            FDTD engine log
      *.fsp               -> <stem>_p0.log next to it
      *.csv/*.xlsx/*.json/*.jsonl   result table (xlsx: every sheet)
      directory           every *_p0.log directly inside it

    options
      --solver auto|fdtd|rcwa   (auto: rcwa if table has Rs/Ts/Rp/Tp and no log given)
      --air-only                FDTD air-only baseline: R+T must be 1
      --lossless                RCWA lossless structure: R+T must be 1 per pol
      --r-col R --t-col T       FDTD R/T column names (default R, T)
      --id-col case_id          case id column (default case_id)
      --log-dir DIR             where <case_id>_p0.log lives (default: table dir and table dir/fsp)
      --sheet NAME              xlsx: only this sheet
      --json OUT                write machine-readable result
      --max-lines N             human summary lines per check (default 8)

Python API:
    from assess_run import assess, parse_p0_log
    res = assess(["fsp/case_p0.log", "results.csv"], air_only=False)
    res["verdict"]  -> "PASS" | "WARN" | "FAIL"
"""
import argparse
import json
import math
import os
import re
import sys

VERDICT_RANK = {"PASS": 0, "WARN": 1, "FAIL": 2}

# Default tolerances. Each one is a CLI option; values are recorded in the output.
TOL_RANGE = 1e-3       # R,T outside [-tol, 1+tol] -> FAIL; outside [0,1] but within tol -> WARN
TOL_AIR = 1e-3         # air-only |R+T-1| above this -> FAIL
TOL_LOSSLESS = 1e-6    # RCWA lossless |R+T-1| above this -> FAIL ("R+T == 1.000000"); also RCWA R+T-1 > this -> FAIL
TOL_SUM_WARN = 1e-3    # FDTD R+T-1 above this -> WARN (normal-incidence 2D slab |1-R-T| <= 8.1e-4)
TOL_SUM_FAIL = 1e-2    # FDTD R+T-1 above this -> FAIL (sound high-angle results <= 0.7 %, errors 3.9 %+)
POL_GAP_WARN = 0.3     # RCWA |eps_s - eps_p| above this -> WARN (fake-plateau check)
SHUTOFF_REL_TOL = 1e-2  # recorded shutoff_time_fs vs log (t_probe last sample lags ~0.1%)


# --------------------------------------------------------------------------- helpers

def _ascii(s):
    return str(s).encode("ascii", "backslashreplace").decode("ascii")


def _to_float(v):
    """Return (value, kind). kind: 'num', 'nonfinite', 'missing', 'text'."""
    if v is None:
        return None, "missing"
    if isinstance(v, bool):
        return float(v), "num"
    if isinstance(v, (int, float)):
        f = float(v)
        return f, ("num" if math.isfinite(f) else "nonfinite")
    if isinstance(v, complex):
        return None, "text"
    s = str(v).strip()
    if s == "":
        return None, "missing"
    try:
        f = float(s)
    except ValueError:
        return None, "text"
    return f, ("num" if math.isfinite(f) else "nonfinite")


class Report:
    def __init__(self):
        self.items = []
        self.notes = []

    def add(self, check, verdict, target, detail, evidence):
        self.items.append({"check": check, "verdict": verdict, "target": target,
                           "detail": detail, "evidence": evidence})

    def verdict(self):
        v = "PASS"
        for it in self.items:
            if VERDICT_RANK[it["verdict"]] > VERDICT_RANK[v]:
                v = it["verdict"]
        return v


# --------------------------------------------------------------------------- FDTD log

# The ONE shutoff/completion parser for this agent.
# Evidence: fdtd_pitfalls_api.md -- the old regex
# "Auto shutoff.*?(\d+)\s*fs" never matched and defaulted every case to 'cap'.
# Completion time is only on "Completed N iterations, or X s of Simulation Time";
# early shutoff is only proven by "autoshutoff criteria are satisfied".
RE_COMPLETED = re.compile(
    r"Completed\s+(\d+)\s+iterations,\s*or\s+([\d.eE+\-]+)\s*s\s+of\s+Simulation\s+Time")
RE_EARLY = "autoshutoff criteria are satisfied"
RE_STARTING = re.compile(r"Starting\s+(\d+)\s+total\s+iterations")
# Evidence: fdtd_pitfalls_api.md -- run() returns
# normally; the only message is e.g. "The program terminated due to an error:Error CPU
# simulation cannot run with FDTD property express mode active."
RE_ENGINE_ERR = re.compile(r"terminated due to an error|^\s*Error\b|\bError:", re.I | re.M)
RE_VERSION = re.compile(r"FDTD Solver Version\s+(\S+)")
RE_FSP = re.compile(r"Running fsp file:.*\s(\S.*\.fsp)\s*$", re.M)


def parse_p0_log(path):
    """Parse one <fsp_stem>_p0.log. Pure function, no verdicts."""
    r = {"log": path, "exists": os.path.exists(path)}
    if not r["exists"]:
        return r
    with open(path, "rb") as f:
        text = f.read().decode("utf-8", errors="replace")
    errs = [m.group(0) for m in RE_ENGINE_ERR.finditer(text)]
    err_lines = []
    for ln in text.splitlines():
        if RE_ENGINE_ERR.search(ln):
            err_lines.append(ln.strip())
    r["engine_errors"] = err_lines[:6] if errs else []
    comps = RE_COMPLETED.findall(text)
    r["n_completed_blocks"] = len(comps)
    if comps:
        n_it, t_s = comps[-1]
        r["completed"] = True
        r["n_iterations"] = int(n_it)
        r["sim_time_fs"] = float(t_s) * 1e15
    else:
        r["completed"] = False
        r["n_iterations"] = None
        r["sim_time_fs"] = None
    st = RE_STARTING.findall(text)
    r["planned_iterations"] = int(st[-1]) if st else None
    r["early_shutoff"] = RE_EARLY in text.lower()
    r["hit_cap"] = (r["completed"] and not r["early_shutoff"])
    m = RE_VERSION.search(text)
    r["solver_version"] = m.group(1) if m else None
    m = RE_FSP.search(text)
    r["fsp_in_log"] = m.group(1).strip() if m else None
    return r


def judge_log(rep, info, label):
    """FDTD log verdicts."""
    if not info["exists"]:
        # Evidence: fdtd_pitfalls_api.md
        # (the _p0.log is the only engine-error channel; no log = unverifiable run)
        rep.add("log_exists", "FAIL", label, "no _p0.log",
                "fdtd_pitfalls_api.md")
        return
    if info["engine_errors"]:
        # Evidence: fdtd_pitfalls_api.md
        rep.add("engine_error", "FAIL", label, " | ".join(info["engine_errors"])[:300],
                "fdtd_pitfalls_api.md")
    else:
        rep.add("engine_error", "PASS", label, "", "")
    if not info["completed"]:
        # Evidence: fdtd_pitfalls_api.md (completion only on the
        # "Completed N iterations" line); missing = engine refused, killed or still running
        rep.add("run_completed", "FAIL", label,
                "no 'Completed N iterations' line (killed / refused / unfinished)",
                "fdtd_pitfalls_api.md")
        return
    rep.add("run_completed", "PASS", label, "", "")
    if info["hit_cap"]:
        # Evidence: fdtd_pitfalls_api.md
        rep.add("autoshutoff", "WARN", label,
                "autoshutoff criteria NOT satisfied; ran to sim time cap (%.1f fs)"
                % info["sim_time_fs"], "fdtd_pitfalls_api.md")
    else:
        rep.add("autoshutoff", "PASS", label,
                "early shutoff at %.1f fs" % info["sim_time_fs"], "")


# --------------------------------------------------------------------------- tables

def load_tables(path, sheet=None):
    """Return list of (label, rows[list of dict]). Values kept raw (no NaN filling)."""
    ext = os.path.splitext(path)[1].lower()
    base = os.path.basename(path)
    if ext == ".csv":
        import csv
        with open(path, newline="", encoding="utf-8-sig") as f:
            lines = [ln for ln in f if not ln.startswith("#")]
        return [(base, list(csv.DictReader(lines)))]
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        out = []
        for name in wb.sheetnames:
            if sheet is not None and name != sheet:
                continue
            ws = wb[name]
            it = ws.iter_rows(values_only=True)
            try:
                hdr = next(it)
            except StopIteration:
                continue
            hdr = [str(h) if h is not None else "" for h in hdr]
            rows = [dict(zip(hdr, r)) for r in it if any(c is not None for c in r)]
            out.append(("%s[%s]" % (base, name), rows))
        wb.close()
        return out
    if ext == ".jsonl":
        with open(path, encoding="utf-8") as f:
            rows = [json.loads(ln) for ln in f if ln.strip()]
        return [(base, rows)]
    if ext == ".json":
        with open(path, encoding="utf-8") as f:
            d = json.load(f)
        if isinstance(d, list):
            return [(base, d)]
        if isinstance(d, dict):
            if d and all(isinstance(v, list) for v in d.values()):
                n = max(len(v) for v in d.values())
                rows = [{k: (v[i] if i < len(v) else None) for k, v in d.items()}
                        for i in range(n)]
                return [(base, rows)]
            return [(base, [d])]
    raise ValueError("unsupported table: %s" % path)


def _row_label(tlabel, i, row, id_col):
    cid = row.get(id_col)
    return "%s:%s" % (tlabel, cid if cid not in (None, "") else "row%d" % (i + 1))


def check_range(rep, label, name, v):
    """0 <= R,T <= 1. Evidence: fdtd_detail_rt_signs.md (sign
    flips give negative R/T) and fdtd_detail_bfast.md (T > 1)."""
    ev = "fdtd_detail_rt_signs.md; fdtd_detail_bfast.md"
    dev = max(-v, v - 1.0, 0.0)
    if dev > TOL_RANGE:
        rep.add("rt_range", "FAIL", label, "%s=%.6g outside [0,1] by %.3g" % (name, v, dev), ev)
    elif dev > 0:
        rep.add("rt_range", "WARN", label, "%s=%.6g outside [0,1] by %.3g (<= tol %.0e)"
                % (name, v, dev, TOL_RANGE), ev)
    else:
        rep.add("rt_range", "PASS", label, "", "")


def check_sum(rep, label, name, R, T, warn, fail, ev):
    """Passive structure: absorptance 1-R-T >= 0, i.e. R+T <= 1. Each of R, T inside [0,1] is not
    enough (review 2026-10-22: Rs=Ts=Rp=Tp=0.8 passed)."""
    excess = R + T - 1.0
    if excess > fail:
        rep.add("rt_sum", "FAIL", label, "%s=%.6f > 1 (1-R-T=%.3g < 0)" % (name, R + T, -excess), ev)
    elif excess > warn:
        rep.add("rt_sum", "WARN", label, "%s=%.6f > 1 by %.3g (<= %.0e)" % (name, R + T, excess, fail), ev)
    else:
        rep.add("rt_sum", "PASS", label, "", "")


def _get_num(rep, label, row, col):
    """Fetch a required numeric R/T cell. Non-finite or missing -> FAIL (never filled).
    Evidence: agents/lumerical.md execution step 4 + lumerical_rcwa.md s6
    (silent NaN R/T)."""
    v, kind = _to_float(row.get(col))
    if kind == "num":
        return v
    rep.add("nan", "FAIL", label, "%s is %s (%r)" % (col, kind, row.get(col)),
            "lumerical_rcwa.md s6")
    return None


def check_ids(rep, tlabel, rows, id_col):
    """case_id uniqueness. Evidence: fdtd_pitfalls_api.md
    (f"a{a:.1f}" mapped 0.15/0.10/0.05 onto one id -> shared .fsp/_p0.log)."""
    ids = [r.get(id_col) for r in rows if r.get(id_col) not in (None, "")]
    if not ids:
        return
    seen, dup = set(), set()
    for c in ids:
        (dup if c in seen else seen).add(c)
    if dup:
        rep.add("case_id_unique", "FAIL", tlabel,
                "duplicate %s: %s" % (id_col, ", ".join(sorted(map(str, dup)))[:300]),
                "fdtd_pitfalls_api.md")
    else:
        rep.add("case_id_unique", "PASS", tlabel, "%d unique" % len(ids), "")


def _log_for_row(row, id_col, table_path, log_dirs):
    fsp = row.get("fsp")
    if fsp in (None, ""):
        fsp = row.get("fsp_path")
    if fsp not in (None, ""):
        fsp = str(fsp)
        cands = [fsp, os.path.join(os.path.dirname(table_path), fsp)]
        cands += [os.path.join(d, os.path.basename(fsp)) for d in log_dirs]
        for c in cands:
            lg = os.path.splitext(c)[0] + "_p0.log"
            if os.path.exists(lg):
                return os.path.normpath(lg)
    cid = row.get(id_col)
    if cid not in (None, ""):
        for d in log_dirs:
            lg = os.path.join(d, "%s_p0.log" % cid)
            if os.path.exists(lg):
                return os.path.normpath(lg)
    return None


def check_recorded_shutoff(rep, label, row, info):
    """Recorded hit_cap / shutoff_time_fs must equal the log.
    Evidence: fdtd_pitfalls_api.md (all 16 cases recorded 'cap'
    while logs showed early shutoff)."""
    ev = "fdtd_pitfalls_api.md"
    if not info.get("completed"):
        return
    if "hit_cap" in row and row["hit_cap"] not in (None, ""):
        rv = row["hit_cap"]
        rec = str(rv).strip().lower() in ("1", "1.0", "true", "yes", "cap")
        if rec != info["hit_cap"]:
            rep.add("shutoff_recorded", "FAIL", label,
                    "table hit_cap=%r but log says %s" % (
                        rv, "cap" if info["hit_cap"] else "early shutoff"), ev)
        else:
            rep.add("shutoff_recorded", "PASS", label, "", "")
    if "shutoff_time_fs" in row and row["shutoff_time_fs"] not in (None, ""):
        rv = row["shutoff_time_fs"]
        v, kind = _to_float(rv)
        if kind == "text":
            rec_cap = str(rv).strip().lower() == "cap"
            if rec_cap != info["hit_cap"]:
                rep.add("shutoff_recorded", "FAIL", label,
                        "table shutoff_time_fs=%r but log says %s at %.1f fs" % (
                            rv, "cap" if info["hit_cap"] else "early shutoff",
                            info["sim_time_fs"]), ev)
        elif kind == "num":
            ref = info["sim_time_fs"]
            if abs(v - ref) > SHUTOFF_REL_TOL * max(abs(ref), 1.0):
                rep.add("shutoff_recorded", "FAIL", label,
                        "table shutoff_time_fs=%.3f vs log %.3f" % (v, ref), ev)
            else:
                rep.add("shutoff_recorded", "PASS", label, "", "")


def assess_fdtd_table(rep, tpath, tlabel, rows, opts, log_cache):
    r_col, t_col, id_col = opts["r_col"], opts["t_col"], opts["id_col"]
    cols = set()
    for r in rows:
        cols.update(r.keys())
    check_ids(rep, tlabel, rows, id_col)
    has_r, has_t = r_col in cols, t_col in cols
    if not (has_r or has_t):
        rep.notes.append("%s: no '%s'/'%s' column, R/T checks skipped" % (tlabel, r_col, t_col))
    if opts["air_only"] and not (has_r and has_t):
        rep.add("air_only_rt", "FAIL", tlabel,
                "--air-only needs both '%s' and '%s'" % (r_col, t_col),
                "fdtd_detail_bfast.md")
    log_dirs = opts["log_dirs"] or [os.path.dirname(tpath),
                                    os.path.join(os.path.dirname(tpath), "fsp")]
    log_owner = {}
    n_linked = 0
    for i, row in enumerate(rows):
        label = _row_label(tlabel, i, row, id_col)
        vals = {}
        for col in ((r_col,) if has_r else ()) + ((t_col,) if has_t else ()):
            vals[col] = _get_num(rep, label, row, col)
            if vals[col] is not None:
                check_range(rep, label, col, vals[col])
        if has_r and has_t and None not in vals.values():
            check_sum(rep, label, "%s+%s" % (r_col, t_col), vals[r_col], vals[t_col], TOL_SUM_WARN, TOL_SUM_FAIL,
                      "fdtd_detail_rt_signs.md; fdtd_detail_bfast.md")
        if opts["air_only"] and has_r and has_t and None not in vals.values():
            # Evidence: fdtd_detail_bfast.md (air-only T=1.039..1.050)
            # and fdtd_detail_pml.md (air-only residual)
            s = vals[r_col] + vals[t_col]
            if abs(s - 1.0) > TOL_AIR:
                rep.add("air_only_rt", "FAIL", label,
                        "R+T=%.6f (|R+T-1|=%.3g > %.0e)" % (s, abs(s - 1.0), TOL_AIR),
                        "fdtd_detail_bfast.md; "
                        "fdtd_detail_pml.md")
            else:
                rep.add("air_only_rt", "PASS", label, "R+T=%.6f" % s, "")
        lg = _log_for_row(row, id_col, tpath, log_dirs)
        if lg is None:
            continue
        n_linked += 1
        # Evidence: fdtd_pitfalls_api.md (shared log path)
        log_owner.setdefault(lg, []).append(label)
        if lg not in log_cache:
            log_cache[lg] = parse_p0_log(lg)
        check_recorded_shutoff(rep, label, row, log_cache[lg])
    for lg, owners in log_owner.items():
        if len(owners) > 1:
            rep.add("log_shared", "FAIL", os.path.basename(lg),
                    "%d rows share this log: %s" % (len(owners), ", ".join(owners)[:300]),
                    "fdtd_pitfalls_api.md")
    if not (has_r or has_t):
        pass
    elif rows and n_linked == 0:
        rep.add("log_link", "WARN", tlabel,
                "no row could be linked to a _p0.log (need 'fsp' or '%s' + --log-dir); "
                "engine status of rows unverified" % id_col,
                "fdtd_pitfalls_api.md")
    elif rows and n_linked < len(rows):
        rep.add("log_link", "WARN", tlabel,
                "%d/%d rows have no _p0.log" % (len(rows) - n_linked, len(rows)),
                "fdtd_pitfalls_api.md")
    return list(log_owner.keys())


RCWA_COLS = ("Rs", "Ts", "Rp", "Tp")


def assess_rcwa_table(rep, tlabel, rows, opts):
    """RCWA has no _p0.log (lumerical_rcwa.md s6): verify the
    numbers themselves."""
    id_col = opts["id_col"]
    cols = set()
    for r in rows:
        cols.update(r.keys())
    check_ids(rep, tlabel, rows, id_col)
    have = [c for c in RCWA_COLS if c in cols]
    if not have:
        rep.notes.append("%s: no Rs/Ts/Rp/Tp columns, skipped" % tlabel)
        return
    pols = [p for p in ("s", "p") if "R" + p in cols and "T" + p in cols]
    if len(pols) < 2:
        # Evidence: lumerical_rcwa.md s4
        rep.add("pol_consistency", "WARN", tlabel,
                "only pol %s present; cross-pol check impossible" % (pols or "none"),
                "lumerical_rcwa.md s4")
    if opts["lossless"] and not pols:
        rep.add("lossless_rt", "FAIL", tlabel, "--lossless needs R/T pairs",
                "lumerical_rcwa.md s3")
    for i, row in enumerate(rows):
        label = _row_label(tlabel, i, row, id_col)
        v = {}
        for c in have:
            # Evidence: lumerical_rcwa.md s6 (Al 15% fill ku=67:
            # Rs/Ts/Rp/Tp all NaN, no exception, no log)
            v[c] = _get_num(rep, label, row, c)
            if v[c] is not None:
                check_range(rep, label, c, v[c])
        eps = {}
        for p in pols:
            R, T = v.get("R" + p), v.get("T" + p)
            if R is None or T is None:
                continue
            eps[p] = 1.0 - R - T
            # Evidence: lumerical_rcwa.md s6 (eps = 1-R-T in [0,1]); v242 matches TMM to 8 digits
            check_sum(rep, label, "R%s+T%s" % (p, p), R, T, TOL_LOSSLESS, TOL_LOSSLESS, "lumerical_rcwa.md s6")
            if opts["lossless"]:
                # Evidence: lumerical_rcwa.md s6 /
                # lumerical_rcwa.md s3 (lossless -> R+T == 1.000000)
                d = abs(R + T - 1.0)
                if d > TOL_LOSSLESS:
                    rep.add("lossless_rt", "FAIL", label,
                            "R%s+T%s=%.9f (|dev| %.3g > %.0e)" % (p, p, R + T, d, TOL_LOSSLESS),
                            "lumerical_rcwa.md s3")
                else:
                    rep.add("lossless_rt", "PASS", label, "", "")
        if len(eps) == 2:
            # Evidence: lumerical_rcwa.md s4
            # (fake plateau: s/TE 0.03 flat vs p/TM 0.84 on the same structure)
            g = abs(eps["s"] - eps["p"])
            if g > POL_GAP_WARN:
                rep.add("pol_consistency", "WARN", label,
                        "eps_s=%.4f eps_p=%.4f gap %.3f > %.2f; judge against analytic"
                        % (eps["s"], eps["p"], g, POL_GAP_WARN),
                        "lumerical_rcwa.md s4")
            else:
                rep.add("pol_consistency", "PASS", label, "", "")


# --------------------------------------------------------------------------- driver

TABLE_EXT = (".csv", ".xlsx", ".json", ".jsonl")


def assess(paths, solver="auto", air_only=False, lossless=False, r_col="R", t_col="T",
           id_col="case_id", log_dirs=None, sheet=None):
    rep = Report()
    logs, tables = [], []
    for p in paths:
        p = os.path.abspath(p)
        if os.path.isdir(p):
            found = sorted(os.path.join(p, f) for f in os.listdir(p) if f.endswith("_p0.log"))
            if not found:  # a wrong folder must not read as "nothing wrong" (review 2026-10-22: empty folder -> PASS)
                rep.add("inputs", "FAIL", p, "no *_p0.log in this folder", "")
            logs += found
        elif p.endswith("_p0.log"):
            logs.append(p)
        elif p.lower().endswith(".fsp"):
            logs.append(os.path.splitext(p)[0] + "_p0.log")
        elif p.lower().endswith(TABLE_EXT):
            tables.append(p)
        else:
            raise ValueError("unsupported path: %s" % p)
    loaded = []
    for t in tables:
        for lab, rows in load_tables(t, sheet):
            loaded.append((t, lab, rows))
    if solver == "auto":
        rc = any(any(c in r for c in RCWA_COLS) for _, _, rows in loaded for r in rows[:1])
        solver = "rcwa" if (rc and not logs) else "fdtd"
    opts = {"air_only": air_only, "lossless": lossless, "r_col": r_col, "t_col": t_col,
            "id_col": id_col, "log_dirs": log_dirs}
    log_cache = {}
    table_logs = []
    for tpath, lab, rows in loaded:
        if solver == "rcwa":
            assess_rcwa_table(rep, lab, rows, opts)
        else:
            table_logs += assess_fdtd_table(rep, tpath, lab, rows, opts, log_cache)
    if loaded and not any(it["check"] in ("rt_range", "nan") for it in rep.items):
        rep.add("rt_columns", "WARN", "-", "no table had usable R/T columns; R/T unchecked", "")
    if solver == "fdtd":
        for lg in logs + table_logs:
            lg = os.path.normpath(lg)
            if lg not in log_cache:
                log_cache[lg] = parse_p0_log(lg)
        for lg in sorted(log_cache):
            judge_log(rep, log_cache[lg], os.path.basename(lg))
    elif logs:
        rep.add("solver", "WARN", "-", "RCWA writes no _p0.log; given logs ignored",
                "lumerical_rcwa.md s6")
    if not rep.items:
        rep.add("inputs", "FAIL", "-", "nothing checked: no _p0.log and no table rows", "")
    return {
        "verdict": rep.verdict(),
        "solver": solver,
        "inputs": {"logs": logs, "tables": tables},
        "options": {"air_only": air_only, "lossless": lossless, "r_col": r_col,
                    "t_col": t_col, "id_col": id_col, "sheet": sheet,
                    "tol_range": TOL_RANGE, "tol_air": TOL_AIR,
                    "tol_lossless": TOL_LOSSLESS, "tol_sum_warn": TOL_SUM_WARN, "tol_sum_fail": TOL_SUM_FAIL,
                    "pol_gap_warn": POL_GAP_WARN},
        "logs": log_cache,
        "checks": rep.items,
        "notes": rep.notes,
    }


def summarize(res, max_lines=8):
    out = []
    out.append("assess_run: %s  (solver=%s, logs=%d, tables=%d)" % (
        res["verdict"], res["solver"], len(res["logs"]), len(res["inputs"]["tables"])))
    by = {}
    for it in res["checks"]:
        by.setdefault(it["check"], {"PASS": 0, "WARN": 0, "FAIL": 0})[it["verdict"]] += 1
    out.append("%-18s %6s %6s %6s" % ("check", "PASS", "WARN", "FAIL"))
    for k, c in by.items():
        out.append("%-18s %6d %6d %6d" % (k, c["PASS"], c["WARN"], c["FAIL"]))
    for level in ("FAIL", "WARN"):
        shown = {}
        for it in res["checks"]:
            if it["verdict"] != level:
                continue
            n = shown.get(it["check"], 0)
            shown[it["check"]] = n + 1
            if n < max_lines:
                out.append("%s %-16s %s  %s  [%s]" % (level, it["check"], it["target"],
                                                       it["detail"], it["evidence"]))
        for k, n in shown.items():
            if n > max_lines:
                out.append("%s %-16s ... %d more" % (level, k, n - max_lines))
    for n in res.get("notes", []):
        out.append("note " + n)
    return _ascii("\n".join(out))


def main(argv=None):
    global TOL_RANGE, TOL_AIR, TOL_LOSSLESS, TOL_SUM_WARN, TOL_SUM_FAIL, POL_GAP_WARN
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("paths", nargs="+")
    ap.add_argument("--solver", default="auto", choices=("auto", "fdtd", "rcwa"))
    ap.add_argument("--air-only", action="store_true")
    ap.add_argument("--lossless", action="store_true")
    ap.add_argument("--r-col", default="R")
    ap.add_argument("--t-col", default="T")
    ap.add_argument("--id-col", default="case_id")
    ap.add_argument("--log-dir", action="append")
    ap.add_argument("--sheet")
    ap.add_argument("--tol-range", type=float, default=TOL_RANGE)
    ap.add_argument("--tol-air", type=float, default=TOL_AIR)
    ap.add_argument("--tol-lossless", type=float, default=TOL_LOSSLESS)
    ap.add_argument("--tol-sum-warn", type=float, default=TOL_SUM_WARN)
    ap.add_argument("--tol-sum-fail", type=float, default=TOL_SUM_FAIL)
    ap.add_argument("--pol-gap-warn", type=float, default=POL_GAP_WARN)
    ap.add_argument("--json")
    ap.add_argument("--max-lines", type=int, default=8)
    a = ap.parse_args(argv)
    TOL_RANGE, TOL_AIR, TOL_LOSSLESS, TOL_SUM_WARN, TOL_SUM_FAIL, POL_GAP_WARN = (
        a.tol_range, a.tol_air, a.tol_lossless, a.tol_sum_warn, a.tol_sum_fail, a.pol_gap_warn)
    res = assess(a.paths, solver=a.solver, air_only=a.air_only, lossless=a.lossless,
                 r_col=a.r_col, t_col=a.t_col, id_col=a.id_col, log_dirs=a.log_dir,
                 sheet=a.sheet)
    print(summarize(res, a.max_lines))
    if a.json:
        with open(a.json, "w", encoding="utf-8") as f:
            json.dump(res, f, indent=1, default=str)
    return {"PASS": 0, "WARN": 1, "FAIL": 2}[res["verdict"]]


if __name__ == "__main__":
    sys.exit(main())
